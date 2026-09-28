# Copyright (c) 2026, NVIDIA CORPORATION.  All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Export per-iteration GPU/host memory metrics from training logs as Excel (xlsx)."""

from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


LOGGER = logging.getLogger(__name__)

MEMORY_LINE_RE = re.compile(r"\(after (\d+) iterations\) memory \(GB\)")

MEMORY_KEYS = (
    "mem-allocated-gigabytes",
    "mem-max-allocated-gigabytes",
    "host-mem-allocated-gigabytes",
    "host-mem-max-active-gigabytes",
)

HEADER = ("iteration", *MEMORY_KEYS)


def parse_memory_rows(log_paths: list[Path]) -> list[list[int | float]]:
    """Parse memory report lines from training logs.

    Each matched line yields one row: the iteration number followed by the
    values of ``MEMORY_KEYS`` in order.

    Args:
        log_paths: Training log files to parse.

    Returns:
        Rows of [iteration, *MEMORY_KEYS values] in log order.
    """
    rows: list[list[int | float]] = []
    for log_path in log_paths:
        with log_path.open(encoding="utf-8", errors="replace") as log_file:
            for line in log_file:
                match = MEMORY_LINE_RE.search(line)
                if match is None:
                    continue
                fields: dict[str, str] = {}
                for part in line[match.end() :].split("|"):
                    key, sep, value = part.strip().partition(":")
                    if sep:
                        fields[key] = value.strip()
                missing = [key for key in MEMORY_KEYS if key not in fields]
                if missing:
                    LOGGER.warning("Skipping memory line missing %s in %s", missing, log_path)
                    continue
                rows.append([int(match.group(1)), *(float(fields[key]) for key in MEMORY_KEYS)])
    return rows


def export_memory(log_paths: list[Path], *, output_path: Path) -> list[list[int | float]]:
    """Extract per-iteration memory metrics and write them as an xlsx table.

    The first row of the sheet is the header (``HEADER``); each subsequent row
    holds one memory report from the logs.

    Args:
        log_paths: Training log files to parse.
        output_path: Destination xlsx file.

    Returns:
        Rows written to the sheet.
    """
    rows = parse_memory_rows(log_paths)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "memory"
    sheet.append(list(HEADER))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for idx, name in enumerate(HEADER, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = max(12, len(name) + 2)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    LOGGER.info("Wrote %d memory rows to %s", len(rows), output_path)
    return rows


def build_parser() -> argparse.ArgumentParser:
    """Build the command-line argument parser."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--log-file",
        action="append",
        required=True,
        type=Path,
        dest="log_files",
        help="Training log to parse; repeat for multiple log files.",
    )
    parser.add_argument("--output", required=True, type=Path, help="Destination xlsx file.")
    return parser


def main() -> int:
    """Run the memory exporter."""
    args = build_parser().parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    export_memory(args.log_files, output_path=args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
