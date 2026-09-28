# Copyright (c) 2026, NVIDIA CORPORATION.  All rights reserved.
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Export per-step DeepSpeed offload training metrics from logs as Excel (xlsx).

Parses lines emitted by
``examples/scale-down/01-analyse/02-deepspeed-offload/train.py`` of the form::

    Step 10: loss=10.998123, time=1.234s, global_tps=5678, \
peak_mem=1.23 GiB, gpu.peak=4.567 GiB, cpu.peak=890.123 MiB

and exports the columns ``step``, ``time`` (s), ``global_tps`` (tokens/s),
``gpu.peak`` (GiB), ``cpu.peak`` (MiB).
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


LOGGER = logging.getLogger(__name__)

STEP_LINE_RE = re.compile(
    r"Step (\S+): loss=\S+, time=([0-9.]+)s, global_tps=([0-9.]+), "
    r"peak_mem=\S+ GiB, gpu\.peak=([0-9.]+) GiB, cpu\.peak=([0-9.]+) MiB"
)

HEADER = ("step", "time", "global_tps", "gpu.peak", "cpu.peak")


def _parse_step(step: str) -> int | str:
    """Return the step as an int when possible, else the raw string."""
    return int(step) if step.isdigit() else step


def parse_step_rows(log_paths: list[Path]) -> list[list[int | str | float]]:
    """Parse per-step metric lines from DeepSpeed offload training logs.

    Each matched line yields one row: step, time (s), global_tps,
    gpu.peak (GiB), cpu.peak (MiB).

    Args:
        log_paths: Training log files to parse.

    Returns:
        Rows of [step, time, global_tps, gpu.peak, cpu.peak] in log order.
    """
    rows: list[list[int | str | float]] = []
    for log_path in log_paths:
        with log_path.open(encoding="utf-8", errors="replace") as log_file:
            for line in log_file:
                match = STEP_LINE_RE.search(line)
                if match is None:
                    continue
                step, time_s, global_tps, gpu_peak, cpu_peak = match.groups()
                rows.append([_parse_step(step), float(time_s), float(global_tps), float(gpu_peak), float(cpu_peak)])
    return rows


def export_deepspeed(log_paths: list[Path], *, output_path: Path) -> list[list[int | str | float]]:
    """Extract per-step metrics and write them as an xlsx table.

    The first row of the sheet is the header (``HEADER``); each subsequent row
    holds one step report from the logs.

    Args:
        log_paths: Training log files to parse.
        output_path: Destination xlsx file.

    Returns:
        Rows written to the sheet.
    """
    rows = parse_step_rows(log_paths)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "deepspeed"
    sheet.append(list(HEADER))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append(row)
    for idx, name in enumerate(HEADER, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = max(12, len(name) + 2)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    LOGGER.info("Wrote %d step rows to %s", len(rows), output_path)
    return rows


def build_parser() -> argparse.ArgumentParser:
    """Build the command-line argument parser."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--log-file",
        action="append",
        required=True,
        type=Path,
        dest="log_files",
        help="Training log to parse; repeat for multiple log files.",
    )
    parser.add_argument("--output", required=True, type=Path, help="Destination xlsx file.")
    return parser


def main() -> int:
    """Run the DeepSpeed metrics exporter."""
    args = build_parser().parse_args()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    export_deepspeed(args.log_files, output_path=args.output)
    return 0


if __name__ == "__main__":
    sys.exit(main())
